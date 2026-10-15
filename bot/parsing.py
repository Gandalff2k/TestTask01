import io
import re

import pymupdf
import pytesseract
from docx import Document as DocxDocument
from PIL import Image, ImageFilter, ImageOps

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")

_MIN_TEXT_CHARS = 20
_OCR_DPI = 400
_OCR_LANG = "ukr"
_OCR_THRESHOLD = 110
MAX_PDF_PAGES = 20
_WORD = re.compile(r"[А-ЩЬЮЯҐЄІЇа-щьюяґєії'’]{4,}")


class UnsupportedFormatError(Exception):
    pass


class EmptyDocumentError(Exception):
    pass


class TooManyPagesError(Exception):
    pass


def is_supported(filename: str) -> bool:
    return filename.lower().endswith(SUPPORTED_EXTENSIONS)


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".txt"):
        text = _from_txt(data)
    elif name.endswith(".docx"):
        text = _from_docx(data)
    elif name.endswith(".pdf"):
        text = _from_pdf(data)
    else:
        raise UnsupportedFormatError(filename)
    text = _normalize(text)
    if not text:
        raise EmptyDocumentError(filename)
    return text


def _from_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _from_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _from_pdf(data: bytes) -> str:
    pages: list[str] = []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        if pdf.page_count > MAX_PDF_PAGES:
            raise TooManyPagesError(pdf.page_count)
        for page in pdf:
            text = page.get_text()
            if len(text.strip()) < _MIN_TEXT_CHARS:
                text = _ocr_page(page)
            pages.append(text)
    return "\n".join(pages)


def _ocr_page(page: "pymupdf.Page") -> str:
    pixmap = page.get_pixmap(dpi=_OCR_DPI)
    image = ImageOps.grayscale(Image.open(io.BytesIO(pixmap.tobytes("png"))))
    image = image.filter(ImageFilter.MedianFilter(3))
    image = image.point(lambda p: 0 if p < _OCR_THRESHOLD else 255)
    text = pytesseract.image_to_string(image, lang=_OCR_LANG)
    return "\n".join(line for line in text.split("\n") if not line.strip() or _WORD.search(line))


def _normalize(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    out: list[str] = []
    for line in lines:
        if not line and (not out or not out[-1]):
            continue
        out.append(line)
    return "\n".join(out).strip()
